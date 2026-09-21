# -*- coding: utf-8 -*-
"""Document parser service — extracts plain text from PDF and DOCX files.

Supports:
- PDF via pdfplumber (full text extraction page by page)
- DOCX via python-docx (paragraphs + table cells)
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

import httpx
import pdfplumber
from docx import Document

from app.core.exceptions import DocumentParsingError, FileDownloadError

logger = logging.getLogger(__name__)


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file.

    Args:
        file_bytes: Raw bytes of the PDF file.

    Returns:
        Concatenated text from all pages.

    Raises:
        DocumentParsingError: If the PDF cannot be parsed.
    """
    try:
        text_parts: list[str] = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(f"--- Página {i + 1} ---\n{page_text}")
                
                # Also extract tables as text
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if row:
                            cells = [str(cell or "").strip() for cell in row]
                            text_parts.append(" | ".join(cells))
        
        result = "\n\n".join(text_parts)
        if not result.strip():
            raise DocumentParsingError(
                message="El PDF no contiene texto extraíble",
                detail="El archivo puede ser un escaneo sin OCR."
            )
        return result
    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(
            message=f"Error al parsear PDF: {e}",
            detail=str(e)
        ) from e


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file.

    Args:
        file_bytes: Raw bytes of the DOCX file.

    Returns:
        Concatenated text from paragraphs and tables.

    Raises:
        DocumentParsingError: If the DOCX cannot be parsed.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_parts: list[str] = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    text_parts.append(" | ".join(cells))

        result = "\n".join(text_parts)
        if not result.strip():
            raise DocumentParsingError(
                message="El archivo DOCX no contiene texto extraíble"
            )
        return result
    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(
            message=f"Error al parsear DOCX: {e}",
            detail=str(e)
        ) from e


def parse_document(filename: str, file_bytes: bytes) -> str:
    """Parse a document based on file extension.

    Args:
        filename: Original filename with extension.
        file_bytes: Raw bytes of the file.

    Returns:
        Extracted plain text.

    Raises:
        DocumentParsingError: If file type is unsupported or parsing fails.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_bytes)
    else:
        raise DocumentParsingError(
            message=f"Tipo de archivo no soportado: {ext}",
            detail="Solo se aceptan archivos PDF y DOCX."
        )


async def download_file(url: str, timeout: float = 60.0) -> tuple[str, bytes]:
    """Download a file from a URL.

    Args:
        url: The URL to download from.
        timeout: Request timeout in seconds.

    Returns:
        Tuple of (filename, file_bytes).

    Raises:
        FileDownloadError: If the download fails.
    """
    try:
        async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
            response = await client.get(url)
            response.raise_for_status()

            # Extract filename from URL or Content-Disposition
            filename = Path(url.split("?")[0]).name
            cd = response.headers.get("content-disposition", "")
            if "filename=" in cd:
                filename = cd.split("filename=")[-1].strip('"\' ')

            if not filename:
                filename = "document.pdf"

            return filename, response.content
    except httpx.HTTPStatusError as e:
        raise FileDownloadError(
            message=f"Error HTTP {e.response.status_code} al descargar {url}",
            detail=str(e)
        ) from e
    except Exception as e:
        raise FileDownloadError(
            message=f"Error al descargar archivo de {url}",
            detail=str(e)
        ) from e
