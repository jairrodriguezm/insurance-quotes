# -*- coding: utf-8 -*-
"""Word (.docx) Renderer for Management Reports (Informes de Gestión).

Generates a pixel-perfect corporate Word document matching
'Informe_Gestion_SONACOL_2026-09.docx' for Multiriesgos de Colombia Ltda.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

# Color palette matching Informe_Gestion_SONACOL_2026-09.docx
HEX_NAVY = "1F3864"
HEX_GRAY_TEXT = "595959"
HEX_FOOTER_GRAY = "7F7F7F"
HEX_BORDER_GRAY = "BFBFBF"
HEX_ZEBRA_FILL = "F2F2F2"
HEX_TOTAL_FILL = "D9E2F3"
HEX_WHITE = "FFFFFF"

RGB_NAVY = RGBColor(0x1F, 0x38, 0x64)
RGB_GRAY_TEXT = RGBColor(0x59, 0x59, 0x59)
RGB_FOOTER_GRAY = RGBColor(0x7F, 0x7F, 0x7F)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RGB_BLACK = RGBColor(0x00, 0x00, 0x00)


def _set_run_font(
    run,
    size_pt: float,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    font_name: str = "Arial",
) -> None:
    """Apply consistent Arial typography to a run."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)


def _set_paragraph_spacing(
    p,
    before_twips: int | None = None,
    after_twips: int | None = None,
    line_twips: int | None = None,
) -> None:
    """Set exact twip spacing on a paragraph (<w:spacing>)."""
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    if before_twips is not None:
        spacing.set(qn("w:before"), str(before_twips))
    if after_twips is not None:
        spacing.set(qn("w:after"), str(after_twips))
    if line_twips is not None:
        spacing.set(qn("w:line"), str(line_twips))


def _set_paragraph_bottom_border(p, color_hex: str = HEX_NAVY, sz: int = 8, space: int = 2) -> None:
    """Add a bottom border to a paragraph (used for Heading 1 and Header)."""
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)


def _add_section_heading(doc: Document, title: str) -> Any:
    """Add a numbered section heading matching Heading 1 in reference docx."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Heading 1"]
    except KeyError:
        pass
    p.paragraph_format.keep_with_next = True
    _set_paragraph_bottom_border(p, color_hex=HEX_NAVY, sz=8, space=2)
    _set_paragraph_spacing(p, before_twips=300, after_twips=120)
    run = p.add_run(title)
    _set_run_font(run, size_pt=13.0, bold=True, color=RGB_NAVY)
    return p


def _add_spacer_paragraph(doc: Document, after_twips: int = 80) -> Any:
    """Add a compact spacer paragraph between tables."""
    p = doc.add_paragraph()
    _set_paragraph_spacing(p, before_twips=0, after_twips=after_twips, line_twips=276)
    return p


def _configure_header_and_footer(doc: Document, assets_dir: str | Path | None = None) -> None:
    """Configure the corporate header (logo + right title + navy rule) and footer (page X of Y)."""
    section = doc.sections[0]
    section.top_margin = Inches(1.0417)
    section.bottom_margin = Inches(0.8333)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

    # Header
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.text = ""
    _set_paragraph_bottom_border(hp, color_hex=HEX_NAVY, sz=6, space=4)
    _set_paragraph_spacing(hp, after_twips=80)

    pPr = hp._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    pPr.append(tabs)

    logo_path: Path | None = None
    if assets_dir:
        candidate = Path(assets_dir) / "logo_mrc.png"
        if candidate.exists():
            logo_path = candidate

    if logo_path and logo_path.exists():
        r_logo = hp.add_run()
        r_logo.add_picture(str(logo_path), width=Inches(1.354), height=Inches(0.479))

    r_title = hp.add_run("\tInforme de gestión de solicitudes")
    _set_run_font(r_title, size_pt=9.0, bold=False, color=RGB_GRAY_TEXT)

    # Footer: "Multiriesgos de Colombia Ltda. · Página X de Y"
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r1 = fp.add_run("Multiriesgos de Colombia Ltda. · Página ")
    _set_run_font(r1, size_pt=8.0, color=RGB_FOOTER_GRAY)

    def _add_xml_field(paragraph, field_name: str) -> None:
        r_fld = paragraph.add_run()
        _set_run_font(r_fld, size_pt=8.0, color=RGB_FOOTER_GRAY)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = field_name
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        r_fld._r.append(fld_begin)
        r_fld._r.append(instr)
        r_fld._r.append(fld_sep)
        r_fld._r.append(fld_end)

    _add_xml_field(fp, "PAGE")
    r2 = fp.add_run(" de ")
    _set_run_font(r2, size_pt=8.0, color=RGB_FOOTER_GRAY)
    _add_xml_field(fp, "NUMPAGES")


def _set_table_grid_and_width(table, col_widths_dxa: list[int]) -> None:
    """Configure exact dxa column widths on tblPr and tblGrid."""
    total_w = sum(col_widths_dxa)
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(total_w))

    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for w in col_widths_dxa:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(w))
        tblGrid.append(gridCol)
    tbl.insert(tbl.index(tblPr) + 1, tblGrid)


def _style_cell_box(
    cell,
    width_dxa: int,
    fill_hex: str | None = None,
    border_color: str = HEX_BORDER_GRAY,
    border_sz: int = 4,
    pad_top: int = 50,
    pad_bottom: int = 50,
    pad_left: int = 90,
    pad_right: int = 90,
    is_kpi_card: bool = False,
) -> None:
    """Apply exact borders, background shading, width, margins, and vertical centering to a cell."""
    tcPr = cell._tc.get_or_add_tcPr()

    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(width_dxa))

    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is not None:
        tcPr.remove(tcBorders)
    if is_kpi_card:
        borders_xml = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none" w:color="FFFFFF" w:sz="0"/>'
            f'<w:left w:val="single" w:color="FFFFFF" w:sz="24"/>'
            f'<w:bottom w:val="none" w:color="FFFFFF" w:sz="0"/>'
            f'<w:right w:val="single" w:color="FFFFFF" w:sz="24"/>'
            f"</w:tcBorders>"
        )
    else:
        borders_xml = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:color="{border_color}" w:sz="{border_sz}"/>'
            f'<w:left w:val="single" w:color="{border_color}" w:sz="{border_sz}"/>'
            f'<w:bottom w:val="single" w:color="{border_color}" w:sz="{border_sz}"/>'
            f'<w:right w:val="single" w:color="{border_color}" w:sz="{border_sz}"/>'
            f"</w:tcBorders>"
        )
    tcPr.append(borders_xml)

    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)
    if fill_hex:
        shd_xml = parse_xml(
            f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>'
        )
        tcPr.append(shd_xml)

    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is not None:
        tcPr.remove(tcMar)
    mar_xml = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:type="dxa" w:w="{pad_top}"/>'
        f'<w:left w:type="dxa" w:w="{pad_left}"/>'
        f'<w:bottom w:type="dxa" w:w="{pad_bottom}"/>'
        f'<w:right w:type="dxa" w:w="{pad_right}"/>'
        f"</w:tcMar>"
    )
    tcPr.append(mar_xml)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _populate_standard_table(
    doc: Document,
    col_widths_dxa: list[int],
    headers: list[str],
    col_alignments: list[WD_ALIGN_PARAGRAPH],
    rows_data: list[list[str]],
    font_size_pt: float = 8.5,
    bold_cols: set[int] | None = None,
    has_total_row: bool = False,
) -> Any:
    """Create and populate a standard corporate table matching the SONACOL report."""
    bold_cols = bold_cols or set()
    total_rows = 1 + len(rows_data)
    table = doc.add_table(rows=total_rows, cols=len(col_widths_dxa))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_grid_and_width(table, col_widths_dxa)

    hdr_cells = table.rows[0].cells
    for c_idx, h_text in enumerate(headers):
        cell = hdr_cells[c_idx]
        _style_cell_box(
            cell,
            width_dxa=col_widths_dxa[c_idx],
            fill_hex=HEX_NAVY,
        )
        p = cell.paragraphs[0]
        p.text = ""
        p.alignment = col_alignments[c_idx] if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        if len(headers) <= 3 and c_idx == len(headers) - 1 and col_alignments[c_idx] == WD_ALIGN_PARAGRAPH.LEFT:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_paragraph_spacing(p, before_twips=30, after_twips=30)
        run = p.add_run(h_text)
        _set_run_font(run, size_pt=font_size_pt, bold=True, color=RGB_WHITE)

    for r_offset, row_vals in enumerate(rows_data):
        r_idx = r_offset + 1
        is_last = r_offset == len(rows_data) - 1
        is_total = has_total_row and is_last

        if is_total:
            fill_hex = HEX_TOTAL_FILL
        elif r_idx % 2 == 0:
            fill_hex = HEX_ZEBRA_FILL
        else:
            fill_hex = None

        row_cells = table.rows[r_idx].cells
        for c_idx, val_text in enumerate(row_vals):
            cell = row_cells[c_idx]
            _style_cell_box(
                cell,
                width_dxa=col_widths_dxa[c_idx],
                fill_hex=fill_hex,
            )
            lines = str(val_text).split("\n")
            for l_idx, line_str in enumerate(lines):
                p = cell.paragraphs[0] if l_idx == 0 else cell.add_paragraph()
                if l_idx == 0:
                    p.text = ""
                p.alignment = col_alignments[c_idx]
                _set_paragraph_spacing(p, before_twips=30, after_twips=30)

                is_bold = is_total or (c_idx in bold_cols and l_idx == 0)
                if l_idx > 0 and line_str.startswith("Integrantes:"):
                    run = p.add_run(line_str)
                    _set_run_font(
                        run,
                        size_pt=max(7.0, font_size_pt - 0.5),
                        bold=False,
                        color=RGB_GRAY_TEXT,
                    )
                else:
                    run = p.add_run(line_str)
                    _set_run_font(
                        run,
                        size_pt=font_size_pt,
                        bold=is_bold,
                        color=RGB_BLACK,
                    )

    return table


def _add_bullet_note(doc: Document, text: str) -> Any:
    """Add a methodological bullet point paragraph matching Section 7 (10pt Arial)."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["List Bullet"]
    except KeyError:
        try:
            p.style = doc.styles["List Paragraph"]
        except KeyError:
            pass
    _set_paragraph_spacing(p, after_twips=60, line_twips=276)
    r_txt = p.add_run(text)
    _set_run_font(r_txt, size_pt=10.0, bold=False, color=RGB_BLACK)
    return p


def render_informe_gestion(
    data: dict[str, Any],
    output_path: str | Path,
    assets_dir: str | Path | None = None,
) -> str:
    """Render the complete Management Report (.docx) from processed Flokzu analytics data."""
    doc = Document()
    _configure_header_and_footer(doc, assets_dir=assets_dir)

    # ------------------------------------------------------------------ #
    # Title Block (Paragraphs 0..3)                                      #
    # ------------------------------------------------------------------ #
    p0 = doc.add_paragraph()
    _set_paragraph_spacing(p0, before_twips=120, after_twips=40)
    r0 = p0.add_run(f"Informe de gestión · {data['periodo_titulo']}")
    _set_run_font(r0, size_pt=11.0, bold=False, color=RGB_GRAY_TEXT)

    p1 = doc.add_paragraph()
    _set_paragraph_spacing(p1, after_twips=40)
    r1 = p1.add_run(data["cliente_nombre_completo"])
    _set_run_font(r1, size_pt=20.0, bold=True, color=RGB_NAVY)

    p2 = doc.add_paragraph()
    _set_paragraph_spacing(p2, before_twips=0, after_twips=60, line_twips=276)
    nit_prefix = f"NIT {data['cliente_nit']} · " if data.get("cliente_nit") else ""
    r2 = p2.add_run(
        f"{nit_prefix}Incluye solicitudes como tomador, asegurado e integrante de consorcios o uniones temporales"
    )
    _set_run_font(r2, size_pt=9.5, bold=False, color=RGB_GRAY_TEXT)

    p3 = doc.add_paragraph()
    _set_paragraph_spacing(p3, before_twips=0, after_twips=200, line_twips=276)
    r3 = p3.add_run(
        f"Ramo: {data['ramo_nombre']} · Solicitudes recibidas {data['rango_fechas_texto']} · "
        f"Estado de las solicitudes al {data['fecha_corte_larga']}"
    )
    _set_run_font(r3, size_pt=9.0, bold=False, color=RGB_GRAY_TEXT)

    # ------------------------------------------------------------------ #
    # 1. Resumen del mes                                                 #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "1. Resumen del mes")

    kpis = data["kpis"]
    kpi_items = [
        (str(kpis["tramitadas"]), "Tramitadas"),
        (str(kpis["emitidas"]), "Emitidas"),
        (str(kpis["novedades"]), "Novedades atendidas"),
        (str(kpis["declinadas"]), "Declinadas"),
        (str(kpis["pendientes"]), "Pendientes"),
    ]
    kpi_widths = [1872, 1872, 1872, 1872, 1872]
    tbl_kpi = doc.add_table(rows=1, cols=5)
    tbl_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_grid_and_width(tbl_kpi, kpi_widths)
    for c_idx, (num_str, label_str) in enumerate(kpi_items):
        cell = tbl_kpi.rows[0].cells[c_idx]
        _style_cell_box(
            cell,
            width_dxa=kpi_widths[c_idx],
            fill_hex=HEX_NAVY,
            pad_top=120,
            pad_bottom=120,
            pad_left=80,
            pad_right=80,
            is_kpi_card=True,
        )
        cp0 = cell.paragraphs[0]
        cp0.text = ""
        cp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(cp0, after_twips=0)
        rn = cp0.add_run(num_str)
        _set_run_font(rn, size_pt=26.0, bold=True, color=RGB_WHITE)

        cp1 = cell.add_paragraph()
        cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(cp1, after_twips=0)
        rl = cp1.add_run(label_str)
        _set_run_font(rl, size_pt=8.5, bold=False, color=RGB_WHITE)

    _add_spacer_paragraph(doc, after_twips=80)

    # Table 1: Indicators (3 columns: 3000, 1500, 4860)
    _populate_standard_table(
        doc,
        col_widths_dxa=[3000, 1500, 4860],
        headers=["Indicador", "Resultado", "Cómo se calcula / alcance"],
        col_alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
        rows_data=[
            [
                "Índice de efectividad de atención",
                kpis["efectividad_str"],
                kpis["efectividad_detalle"],
            ],
            [
                "Índice de gestión",
                kpis["gestion_str"],
                kpis["gestion_detalle"],
            ],
            [
                "Tiempo promedio de atención",
                kpis["tiempo_promedio_str"],
                kpis["tiempo_promedio_detalle"],
            ],
        ],
        font_size_pt=9.0,
        bold_cols={0, 1},
    )

    _add_spacer_paragraph(doc, after_twips=60)

    p_res = doc.add_paragraph()
    _set_paragraph_spacing(p_res, before_twips=60, after_twips=120, line_twips=276)
    r_res = p_res.add_run(data["resumen_narrativa"])
    _set_run_font(r_res, size_pt=10.0, bold=False, color=RGB_BLACK)

    # ------------------------------------------------------------------ #
    # 2. Trazabilidad de las solicitudes                                 #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "2. Trazabilidad de las solicitudes")

    traz_rows_matrix: list[list[str]] = []
    for r in data["trazabilidad_rows"]:
        traz_rows_matrix.append(
            [
                r["tipo"],
                r["tramitadas"],
                r["emitidas"],
                r["novedades"],
                r["declinadas"],
                r["pendientes"],
                r["efectividad"],
                r["dias_prom"],
            ]
        )
    tot = data["trazabilidad_total"]
    traz_rows_matrix.append(
        [
            tot["tipo"],
            tot["tramitadas"],
            tot["emitidas"],
            tot["novedades"],
            tot["declinadas"],
            tot["pendientes"],
            tot["efectividad"],
            tot["dias_prom"],
        ]
    )

    _populate_standard_table(
        doc,
        col_widths_dxa=[1900, 1050, 1000, 1200, 1050, 1050, 1160, 950],
        headers=[
            "Tipo de solicitud",
            "Tramitadas",
            "Emitidas",
            "Novedades atendidas",
            "Declinadas",
            "Pendientes",
            "Efectividad",
            "Días prom.",
        ],
        col_alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
        rows_data=traz_rows_matrix,
        font_size_pt=8.0,
        has_total_row=True,
    )

    p_traz_note = doc.add_paragraph()
    _set_paragraph_spacing(p_traz_note, before_twips=80, after_twips=120, line_twips=276)
    r_traz_note = p_traz_note.add_run(
        "Emitidas: pólizas nuevas expedidas. Novedades atendidas: modificaciones "
        "(prórrogas, adiciones, aclaraciones) ya resueltas. Días promedio: tiempo "
        "transcurrido entre la recepción de la solicitud y su cierre."
    )
    _set_run_font(r_traz_note, size_pt=8.0, bold=False, color=RGB_GRAY_TEXT)

    # ------------------------------------------------------------------ #
    # 3. Solicitudes declinadas y motivo                                 #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "3. Solicitudes declinadas y motivo")

    declinadas_detalle = data["declinadas_detalle_rows"]
    if not declinadas_detalle:
        p_dec = doc.add_paragraph()
        _set_paragraph_spacing(p_dec, before_twips=0, after_twips=120, line_twips=276)
        r_dec = p_dec.add_run(
            f"{data['periodo_narrativa']} no se declinó ninguna solicitud."
        )
        _set_run_font(r_dec, size_pt=10.0, bold=False, color=RGB_BLACK)
    else:
        cnt_dec = len(declinadas_detalle)
        p_dec = doc.add_paragraph()
        _set_paragraph_spacing(p_dec, before_twips=0, after_twips=120, line_twips=276)
        dec_txt = (
            f"{data['periodo_narrativa']} se declinó 1 solicitud, por los siguientes motivos:"
            if cnt_dec == 1
            else f"{data['periodo_narrativa']} se declinaron {cnt_dec} solicitudes, por los siguientes motivos:"
        )
        r_dec = p_dec.add_run(dec_txt)
        _set_run_font(r_dec, size_pt=10.0, bold=False, color=RGB_BLACK)

        _populate_standard_table(
            doc,
            col_widths_dxa=[6360, 3000],
            headers=["Motivo", "Solicitudes"],
            col_alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
            rows_data=[[r["motivo"], r["cantidad"]] for r in data["declinadas_resumen_rows"]],
            font_size_pt=8.5,
        )
        _add_spacer_paragraph(doc, after_twips=80)
        _populate_standard_table(
            doc,
            col_widths_dxa=[1250, 1100, 1800, 2400, 2810],
            headers=["Solicitud", "Recibida", "Tipo", "Tomador / Asegurado", "Motivo"],
            col_alignments=[
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
            ],
            rows_data=[
                [
                    r["id"],
                    r["recibida"],
                    r["tipo"],
                    r["tomador_asegurado"],
                    r["motivo_declinada"],
                ]
                for r in declinadas_detalle
            ],
            font_size_pt=8.0,
        )

    # ------------------------------------------------------------------ #
    # 4. Solicitudes pendientes y motivo                                 #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "4. Solicitudes pendientes y motivo")

    pendientes_detalle = data["pendientes_detalle_rows"]
    if not pendientes_detalle:
        p_pen = doc.add_paragraph()
        _set_paragraph_spacing(p_pen, before_twips=0, after_twips=120, line_twips=276)
        r_pen = p_pen.add_run(
            f"Al {data['fecha_corte_larga']} no queda ninguna solicitud pendiente."
        )
        _set_run_font(r_pen, size_pt=10.0, bold=False, color=RGB_BLACK)
    else:
        cnt_pen = len(pendientes_detalle)
        p_pen = doc.add_paragraph()
        _set_paragraph_spacing(p_pen, before_twips=0, after_twips=120, line_twips=276)
        pen_txt = (
            f"Al {data['fecha_corte_larga']} queda 1 solicitud pendiente, por los siguientes motivos:"
            if cnt_pen == 1
            else f"Al {data['fecha_corte_larga']} quedan {cnt_pen} solicitudes pendientes, por los siguientes motivos:"
        )
        r_pen = p_pen.add_run(pen_txt)
        _set_run_font(r_pen, size_pt=10.0, bold=False, color=RGB_BLACK)

        _populate_standard_table(
            doc,
            col_widths_dxa=[6360, 3000],
            headers=["Motivo", "Solicitudes"],
            col_alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
            rows_data=[[r["motivo"], r["cantidad"]] for r in data["pendientes_resumen_rows"]],
            font_size_pt=8.5,
        )
        _add_spacer_paragraph(doc, after_twips=80)
        _populate_standard_table(
            doc,
            col_widths_dxa=[1250, 1100, 1800, 2400, 2810],
            headers=["Solicitud", "Recibida", "Tipo", "Tomador / Asegurado", "Motivo"],
            col_alignments=[
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
            ],
            rows_data=[
                [
                    r["id"],
                    r["recibida"],
                    r["tipo"],
                    r["tomador_asegurado"],
                    r["motivo_pendiente"],
                ]
                for r in pendientes_detalle
            ],
            font_size_pt=8.0,
        )

    # ------------------------------------------------------------------ #
    # 5. Tiempos de atención                                             #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "5. Tiempos de atención")

    p_t = doc.add_paragraph()
    _set_paragraph_spacing(p_t, before_twips=0, after_twips=120, line_twips=276)
    r_t = p_t.add_run(
        "El tiempo de atención se mide desde que la solicitud ingresa a Multiriesgos de Colombia "
        "hasta que se cierra (póliza emitida, novedad atendida o solicitud declinada). Incluye los "
        "tiempos de espera de terceros, como la respuesta de la aseguradora o la entrega de documentos "
        "por parte del cliente."
    )
    _set_run_font(r_t, size_pt=10.0, bold=False, color=RGB_BLACK)

    tm = data["tiempos_medidas"]
    _populate_standard_table(
        doc,
        col_widths_dxa=[4680, 4680],
        headers=["Medida", "Resultado"],
        col_alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
        rows_data=[
            ["Solicitudes cerradas en el período", tm["cerradas"]],
            ["Promedio", tm["promedio"]],
            ["Mediana", tm["mediana"]],
            ["Más rápida / más lenta", tm["extremos"]],
        ],
        font_size_pt=8.5,
        bold_cols={0},
    )

    _add_spacer_paragraph(doc, after_twips=80)

    _populate_standard_table(
        doc,
        col_widths_dxa=[6360, 3000],
        headers=["Distribución por tiempo de atención", "Solicitudes"],
        col_alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
        rows_data=[[r["tramo"], r["solicitudes"]] for r in data["tiempos_tramos_rows"]],
        font_size_pt=8.5,
    )

    if data["tiempos_por_tipo_rows"]:
        _add_spacer_paragraph(doc, after_twips=80)
        _populate_standard_table(
            doc,
            col_widths_dxa=[6360, 3000],
            headers=["Promedio por tipo de solicitud", "Días"],
            col_alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
            rows_data=[[r["tipo"], r["promedio"]] for r in data["tiempos_por_tipo_rows"]],
            font_size_pt=8.5,
        )

    # ------------------------------------------------------------------ #
    # 6. Detalle de solicitudes                                          #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "6. Detalle de solicitudes")

    _populate_standard_table(
        doc,
        col_widths_dxa=[1150, 1000, 1450, 2200, 1850, 900, 810],
        headers=["Solicitud", "Recibida", "Tipo", "Tomador", "Asegurado", "Estado", "Días"],
        col_alignments=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
        rows_data=[
            [
                r["id"],
                r["recibida"],
                r["tipo_detalle"],
                r["tomador_tabla8"],
                r["asegurado_fmt"],
                r["estado_label"],
                r["dias_fmt"],
            ]
            for r in data["detalle_solicitudes_rows"]
        ],
        font_size_pt=7.5,
    )

    p_det_note = doc.add_paragraph()
    _set_paragraph_spacing(p_det_note, before_twips=80, after_twips=120, line_twips=276)
    r_det_note = p_det_note.add_run(
        "El nombre de las columnas “Tomador” y “Asegurado” corresponde al remitente y a la entidad "
        "registrados en cada solicitud. Cuando el cliente actúa como integrante de un consorcio o "
        "unión temporal, se indican los integrantes registrados."
    )
    _set_run_font(r_det_note, size_pt=8.0, bold=False, color=RGB_GRAY_TEXT)

    # ------------------------------------------------------------------ #
    # 7. Notas metodológicas                                             #
    # ------------------------------------------------------------------ #
    _add_section_heading(doc, "7. Notas metodológicas")

    _add_bullet_note(
        doc,
        f"Universo: solicitudes {data['universo_periodo_texto']} en las que el cliente aparece como "
        f"tomador, como asegurado o como integrante de un consorcio o unión temporal.",
    )
    _add_bullet_note(
        doc,
        f"Corte: el estado de cada solicitud corresponde al {data['fecha_corte_larga']}; las "
        f"solicitudes pendientes pueden cerrarse después de esa fecha.",
    )
    _add_bullet_note(
        doc,
        "Índice de efectividad de atención = (pólizas emitidas + novedades atendidas) ÷ solicitudes tramitadas.",
    )
    _add_bullet_note(
        doc,
        "Índice de gestión = (pólizas emitidas + novedades atendidas + solicitudes declinadas) ÷ solicitudes tramitadas.",
    )
    pagare_cnt = data.get("pending_pagare_cartera_count", 0)
    pagare_suffix = (
        "en este período no hubo solicitudes en esa condición."
        if pagare_cnt == 0
        else f"en este período hubo {pagare_cnt} {'solicitud' if pagare_cnt == 1 else 'solicitudes'} en esa condición."
    )
    _add_bullet_note(
        doc,
        f"Los pagarés pendientes y los estados de cartera se informan como motivo de pendiente cuando "
        f"la solicitud los requiere; {pagare_suffix}",
    )

    output_path_str = str(output_path)
    doc.save(output_path_str)
    logger.info("Informe de gestión generado exitosamente en: %s", output_path_str)
    return output_path_str
