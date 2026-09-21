# -*- coding: utf-8 -*-
"""DOCX Renderer service — generates the Phase II comparative Word document.

Refactored from the proven fase2_word.py script into a service module.
Uses python-docx directly (NOT docxtpl) for maximum control over dynamic tables,
cell merging, embedded logos, and styling.

The renderer accepts a ConsolidatedData dict (datos.json format) and produces
a .docx file using the corporate letterhead template.
"""
from __future__ import annotations

import logging
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.core.config import settings

logger = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Constants                                                                    #
# --------------------------------------------------------------------------- #
AZUL = RGBColor(0x1F, 0x38, 0x64)
AZUL_HEX = "1F3864"
GRIS_HEX = "F2F2F2"
AZUL_CLARO_HEX = "D9E2F3"
FUENTE = "Arial"
RELLENO = "NO ESPECIFICA"
ROMANOS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

# --------------------------------------------------------------------------- #
# Catalog (inline copy to avoid sys.path manipulation)                         #
# --------------------------------------------------------------------------- #
COBERTURAS = [
    ("__seccion__", "Cobertura Sección 1 Todo Riesgo Construcción"),
    ("danos_materiales", "Daños materiales (Cobertura A)"),
    ("terremoto", "Terremoto, temblor y/o erupción volcánica (Cobertura B)"),
    ("tormenta_inundacion", "Tormenta e inundación (Cobertura C)"),
    ("mantenimiento_amplio", "Mantenimiento amplio (Cobertura D)"),
    ("remocion_escombros", "Remoción de escombros (Cobertura G)"),
    ("hmacc_amit", "Huelga, asonada, motín, conmoción civil o popular"),
    ("hurto_calificado", "Hurto calificado"),
    ("__subseccion__", "Coberturas adicionales"),
    ("cronograma_avance", "Cronograma de avance"),
    ("gastos_horas_extra", "Gastos adicionales de horas extra, trabajo nocturno, flete"),
    ("obras_zona_sismica", "Obras en zona sísmica"),
    ("bienes_fuera_sitio", "Bienes almacenados fuera del sitio"),
    ("prueba_maquinaria", "Prueba de maquinarias e instalaciones"),
    ("campamentos_almacenes", "Campamentos y almacenes de materiales de construcción"),
    ("medidas_inundacion", "Medidas de seguridad contra inundaciones"),
    ("proteccion_incendio", "Protección contra incendio"),
    ("transportes_nacionales", "Transportes nacionales"),
    ("siniestros_serie", "Siniestros en serie"),
    ("obras_civiles_operacion", "Obras civiles aseguradas puestas en operación"),
    ("cimentacion_pilotaje", "Cimentación por pilotaje y tablestacados para fosas de obras"),
    ("hundimiento_subsuelo", "Exclusión de daños y pérdidas por hundimiento y asentamiento del subsuelo"),
    ("error_diseno", "Error de diseño"),
    ("honorarios_profesionales", "Honorarios profesionales"),
    ("planos_documentos", "Planos y documentos"),
    ("gastos_extincion", "Gastos de extinción del siniestro"),
    ("condiciones_remocion", "Condiciones especiales relativas a la remoción de escombros"),
    ("actos_autoridad", "Actos de autoridad"),
    ("preservacion_bienes", "Gastos para la preservación de bienes en caso de pérdida"),
    ("__seccion__", "Cobertura Sección 2 Responsabilidad Civil Extracontractual"),
    ("rce", "Responsabilidad civil extracontractual"),
    ("rce_contratistas", "Contratistas y subcontratistas independientes"),
    ("rce_patronal", "Civil Patronal"),
    ("rce_cruzada", "Civil cruzada"),
    ("rce_vehiculos", "Vehículos propios y no propios"),
    ("rce_cuidado_control", "Bienes bajo cuidado, tenencia y control"),
    ("rce_contaminacion", "Contaminación, polución, filtración, accidental, súbita e imprevista"),
    ("rce_vibracion", "Vibración, eliminación o debilitación de elementos portantes para actividades de construcción"),
    ("rce_subterraneas", "Amparo opcional de conducciones y/o servicios subterráneos"),
    ("__seccion__", "Cobertura Sección 3 Terrorismo"),
    ("terrorismo", "Terrorismo"),
]

DEDUCIBLES = [
    ("__seccion__", "Cobertura Sección 1 Todo Riesgo Construcción", ""),
    ("ded_incendio", "Incendio, rayo, humo y explosión", "10% valor pérdida - 5 SMMLV"),
    ("ded_terremoto", "Terremoto, temblor y/o erupción volcánica", "3% valor del proyecto - 4 SMMLV"),
    ("ded_tormenta", "Tormenta e inundación, huracán, tifón, ciclón", "20% valor de la pérdida - 7 SMMLV"),
    ("ded_mantenimiento", "Mantenimiento amplio", ""),
    ("ded_remocion", "Remoción de escombros", ""),
    ("ded_hmacc", "Huelga, asonada, motín, conmoción civil o popular", "15% del valor de la pérdida - 7 SMMLV"),
    ("ded_hurto", "Hurto calificado", "20% valor pérdida - 7 SMMLV"),
    ("ded_error_diseno", "Error de diseño", "20% valor pérdida - 15 SMMLV"),
    ("ded_cables", "Cables y tuberías subterráneas", "15% valor de la pérdida - 5 SMMLV"),
    ("ded_adyacentes", "Propiedades adyacentes", "20% valor pérdida - 7 SMMLV"),
    ("ded_campamentos", "Campamentos y almacenes", "Aplica deducible amparo afectado"),
    ("ded_hundimiento", "Hundimiento de tierra o desprendimiento de tierra o de roca", "20% valor de la pérdida - 7 SMMLV"),
    ("ded_aeronaves", "Caída de aeronaves", "10% valor de la pérdida - 5 SMMLV"),
    ("ded_impericia", "Impericia, negligencia y actos individuales malintencionados de operadores y trabajadores del asegurado", "10% valor de la pérdida - 5 SMMLV"),
    ("ded_corto_circuito", "Corto circuito", ""),
    ("__seccion__", "Cobertura Sección 2 Responsabilidad Civil Extracontractual", ""),
    ("ded_rce", "Responsabilidad civil extracontractual (E y F)", "10% pérdida - 1 SMMLV"),
    ("ded_rce_cruzada", "Responsabilidad civil cruzada", "15% pérdida - 5 SMMLV"),
    ("ded_propiedades_existentes", "Propiedades existentes o que quedan bajo cuidado, custodia o supervisión del asegurado", ""),
]


# --------------------------------------------------------------------------- #
# Utility functions (from original fase2_word.py)                              #
# --------------------------------------------------------------------------- #

def _fmt_cop(valor: Any, decimales: bool = True) -> str:
    """Format a value as Colombian pesos or return as-is if string."""
    if valor is None:
        return RELLENO
    if isinstance(valor, str):
        return valor
    entero = f"{valor:,.2f}" if decimales else f"{valor:,.0f}"
    entero = entero.replace(",", "@").replace(".", ",").replace("@", ".")
    return f"$ {entero}"


def _valor(aseguradora: dict, clave: str, bloque: str = "coberturas") -> Any:
    """Get a value from an insurer's coverage or deductible block."""
    dato = (aseguradora.get(bloque) or {}).get(clave, None)
    return dato if dato not in (None, "") else RELLENO


def _etiqueta_deducible(etiqueta: str, referencia: str) -> str:
    """Format a deductible label with its market reference."""
    return f"{etiqueta} ({referencia})" if referencia else etiqueta


def _sombrear(celda: Any, hex_color: str) -> None:
    """Apply background shading to a cell."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    celda._tc.get_or_add_tcPr().append(shd)


def _bordes_tabla(tabla: Any, color: str = "7F7F7F", sz: int = 4) -> None:
    """Apply borders to a table."""
    tblPr = tabla._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _repetir_encabezado(fila: Any) -> None:
    """Mark a table row to repeat as header on page breaks."""
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_partir(fila: Any) -> None:
    """Prevent a table row from splitting across pages."""
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _texto(
    celda: Any,
    contenido: str,
    *,
    negrita: bool = False,
    tam: int = 9,
    color: RGBColor | None = None,
    centrado: bool = True,
    fuente: str = FUENTE,
) -> None:
    """Write formatted text to a table cell, handling multi-line content."""
    celda.text = ""
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    for i, linea in enumerate(str(contenido).split("\n")):
        if i:
            p = celda.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
        run = p.add_run(linea)
        run.font.name = fuente
        run.font.size = Pt(tam)
        run.bold = negrita
        if color is not None:
            run.font.color.rgb = color
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _nombre_corto(ase: dict) -> str:
    """Return a clean short display name for the insurer."""
    nombre = ase.get("nombre", "")
    id_ase = ase.get("id", "").lower()
    mapeo = {
        "sura": "SURA",
        "suramericana": "SURA",
        "zurich": "Zurich",
        "chubb": "Chubb",
        "axa": "AXA Colpatria",
        "axa_colpatria": "AXA Colpatria",
        "colpatria": "AXA Colpatria",
        "mundial": "Seguros Mundial",
        "hdi": "HDI",
        "bolivar": "Bolívar",
        "estado": "Seguros del Estado",
        "berkley": "Berkley",
        "davivienda": "Seguros Davivienda",
    }
    for k, v in mapeo.items():
        if k == id_ase or k in id_ase or k in nombre.lower():
            return v
    return nombre.split(" S.A")[0].strip() or id_ase.upper()


def _logo_en_celda(
    celda: Any,
    ruta: str,
    ancho_in: float = 1.25,
    pie: str | None = None,
    texto_alternativo: str | None = None,
) -> None:
    """Insert a logo image into a table cell, falling back to text if missing."""
    celda.text = ""
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ruta and os.path.exists(ruta):
        p.add_run().add_picture(ruta, width=Inches(ancho_in))
    elif texto_alternativo:
        r = p.add_run(texto_alternativo)
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = FUENTE
        r.font.color.rgb = AZUL
    if pie:
        p2 = celda.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(pie)
        r.bold = True
        r.font.size = Pt(8)
        r.font.name = FUENTE
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _titulo(doc: Document, indice: int, etiqueta: str) -> Any:
    """Add a section title with Roman numeral."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{ROMANOS[indice]}.   {etiqueta}")
    run.bold = True
    run.underline = True
    run.font.size = Pt(12)
    run.font.name = FUENTE
    run.font.color.rgb = AZUL
    return p


def _parrafo(
    doc: Document,
    contenido: str,
    *,
    tam: int = 10,
    negrita: bool = False,
    centrado: bool = False,
    vineta: bool = False,
) -> Any:
    """Add a formatted paragraph to the document."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    if vineta:
        p.paragraph_format.left_indent = Inches(0.3)
        contenido = "•  " + contenido
    run = p.add_run(contenido)
    run.font.size = Pt(tam)
    run.font.name = FUENTE
    run.bold = negrita
    return p


def _fila_banda(tabla: Any, texto_banda: str, n_cols: int, principal: bool = True) -> Any:
    """Add a section/subsection header row spanning all columns."""
    fila = tabla.add_row()
    celda = fila.cells[0]
    for otra in fila.cells[1:]:
        celda = celda.merge(otra)
    color = RGBColor(0xFF, 0xFF, 0xFF) if principal else AZUL
    _texto(celda, texto_banda, negrita=True, tam=9, color=color)
    _sombrear(celda, AZUL_HEX if principal else AZUL_CLARO_HEX)
    return fila


# --------------------------------------------------------------------------- #
# Document sections                                                            #
# --------------------------------------------------------------------------- #

def _seccion_general(doc: Document, meta: dict, idx: int) -> None:
    """Render Section I: General Information."""
    _titulo(doc, idx, "INFORMACIÓN GENERAL")
    vig: list[str] = []
    vc = meta.get("vigencia_construccion") or {}
    vm = meta.get("vigencia_mantenimiento") or {}
    if vc:
        vig.append(
            f"Periodo de construcción:\nDesde: {vc.get('desde', '')}   "
            f"Hasta: {vc.get('hasta', '')}"
        )
    if vm:
        vig.append(
            f"Periodo de mantenimiento ({vm.get('tipo', '')} – {vm.get('duracion', '')}):\n"
            f"Desde: {vm.get('desde', '')}   Hasta: {vm.get('hasta', '')}"
        )
    filas = [
        ("FECHA", meta.get("fecha", "")),
        ("TIPO DE COBERTURA", meta.get("tipo_cobertura", "")),
        ("TOMADOR", meta.get("tomador", "")),
        ("ASEGURADO", meta.get("asegurado", "")),
        ("BENEFICIARIO", meta.get("beneficiario", "")),
        ("VIGENCIA", "\n\n".join(vig)),
        ("UBICACIÓN", meta.get("ubicacion", "")),
        ("VALOR ASEGURADO", _fmt_cop(meta.get("valor_asegurado"))),
        ("DESCRIPCIÓN DEL PROYECTO", meta.get("descripcion_proyecto", "")),
    ]
    tabla = doc.add_table(rows=0, cols=2)
    _bordes_tabla(tabla)
    for etq, val in filas:
        fila = tabla.add_row()
        fila.cells[0].width = Inches(1.7)
        fila.cells[1].width = Inches(4.8)
        _texto(fila.cells[0], etq, negrita=True, tam=9, centrado=False)
        _sombrear(fila.cells[0], GRIS_HEX)
        _texto(fila.cells[1], val, tam=9, centrado=False)


def _seccion_cotizaciones(doc: Document, datos: dict, assets: str, idx: int) -> None:
    """Render Section II: Quotes Received (dynamic table with logos)."""
    _titulo(doc, idx, "COTIZACIONES REALIZADAS")
    tabla = doc.add_table(rows=1, cols=4)
    _bordes_tabla(tabla)
    encabezados = ["COMPAÑÍA DE SEGUROS", "TASA", "PRIMA", "MODALIDAD DE ASEGURAMIENTO"]
    anchos = [1.9, 1.0, 1.6, 2.0]
    for celda, etq, ancho in zip(tabla.rows[0].cells, encabezados, anchos):
        _texto(celda, etq, negrita=True, tam=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        _sombrear(celda, AZUL_HEX)
        celda.width = Inches(ancho)

    for ase in datos["aseguradoras"]:
        opciones = ase.get("opciones", [])
        primera = None
        for i, op in enumerate(opciones):
            fila = tabla.add_row()
            for celda, ancho in zip(fila.cells, anchos):
                celda.width = Inches(ancho)
            if i == 0:
                primera = fila.cells[0]
                celda_logo = fila.cells[0]
                celda_logo.text = ""
                cos = ase.get("coaseguro") or [{"logo": ase.get("logo"), "participacion": ""}]
                for j, co in enumerate(cos):
                    p = celda_logo.paragraphs[0] if j == 0 else celda_logo.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    ruta = os.path.join(assets, "logos", co.get("logo") or "")
                    ancho_logo = 1.3 if len(cos) <= 1 else 1.0
                    if os.path.exists(ruta):
                        p.add_run().add_picture(ruta, width=Inches(ancho_logo))
                    else:
                        p.add_run(co.get("nombre", ase.get("nombre", ""))).bold = True
                    if co.get("participacion"):
                        pp = celda_logo.add_paragraph()
                        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r = pp.add_run(co["participacion"])
                        r.bold = True
                        r.font.size = Pt(8)
                        r.font.name = FUENTE
                celda_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            short_name = _nombre_corto(ase)
            etq = op.get("etiqueta", "")
            if short_name and f"({short_name})" not in etq:
                etq_display = f"{etq} ({short_name})"
            else:
                etq_display = etq
            _texto(fila.cells[1], f"{etq_display}\n{op.get('tasa', '')}", tam=9)
            _texto(fila.cells[2], _fmt_cop(op.get("prima")), tam=9, negrita=True)
            _texto(fila.cells[3], op.get("modalidad", ""), tam=9)
            _no_partir(fila)
        if primera is not None and len(opciones) > 1:
            ultima = tabla.rows[-1].cells[0]
            primera.merge(ultima)


def _seccion_recomendacion(doc: Document, datos: dict, assets: str, idx: int) -> None:
    """Render Section III: Recommendation."""
    rec = datos.get("recomendacion") or {}
    if not rec:
        return
    _titulo(doc, idx, "RECOMENDACIÓN")
    elegida = next(
        (a for a in datos["aseguradoras"] if a["id"] == rec.get("aseguradora_id")),
        None,
    )
    tabla = doc.add_table(rows=1, cols=2)
    _bordes_tabla(tabla)
    izq, der = tabla.rows[0].cells
    izq.width = Inches(2.0)
    der.width = Inches(4.5)
    _sombrear(izq, AZUL_CLARO_HEX)
    if elegida:
        pie_rec = rec.get("opcion", "")
        short_rec = _nombre_corto(elegida)
        if short_rec and f"({short_rec})" not in pie_rec:
            pie_rec = f"{pie_rec} ({short_rec})"
        _logo_en_celda(
            izq,
            os.path.join(assets, "logos", elegida.get("logo", "")),
            1.5,
            pie=pie_rec,
            texto_alternativo=short_rec,
        )
    else:
        _texto(izq, rec.get("opcion", ""), negrita=True)
    der.text = ""
    for i, vin in enumerate(rec.get("vinetas", [])):
        p = der.paragraphs[0] if i == 0 else der.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("•  " + vin)
        r.font.size = Pt(9)
        r.font.name = FUENTE
    der.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _tabla_matriz(
    doc: Document,
    filas_catalogo: list,
    aseguradoras: list[dict],
    bloque: str,
    assets: str,
    ancho_concepto: float = 1.9,
) -> Any:
    """Render a comparative matrix table (coverages or deductibles)."""
    n = len(aseguradoras)
    if n == 0:
        return None
    ancho_ase = (6.5 - ancho_concepto) / n
    tabla = doc.add_table(rows=1, cols=n + 1)
    _bordes_tabla(tabla)
    tabla.rows[0].cells[0].width = Inches(ancho_concepto)
    _sombrear(tabla.rows[0].cells[0], AZUL_HEX)
    for i, ase in enumerate(aseguradoras):
        celda = tabla.rows[0].cells[i + 1]
        celda.width = Inches(ancho_ase)
        _logo_en_celda(
            celda,
            os.path.join(assets, "logos", ase.get("logo", "")),
            min(1.2, ancho_ase - 0.15),
            texto_alternativo=_nombre_corto(ase),
        )
    _repetir_encabezado(tabla.rows[0])

    for item in filas_catalogo:
        clave, etiqueta = item[0], item[1]
        referencia = item[2] if len(item) > 2 else ""
        if clave in ("__seccion__", "__subseccion__"):
            _fila_banda(tabla, etiqueta, n + 1, principal=(clave == "__seccion__"))
            continue
        fila = tabla.add_row()
        fila.cells[0].width = Inches(ancho_concepto)
        etq = _etiqueta_deducible(etiqueta, referencia) if bloque == "deducibles" else etiqueta
        _texto(fila.cells[0], etq, tam=8, centrado=False, negrita=False)
        _sombrear(fila.cells[0], GRIS_HEX)
        for i, ase in enumerate(aseguradoras):
            celda = fila.cells[i + 1]
            celda.width = Inches(ancho_ase)
            dato = _valor(ase, clave, bloque)
            _texto(celda, _fmt_cop(dato) if isinstance(dato, (int, float)) else dato, tam=8)
        _no_partir(fila)
    return tabla


def _seccion_coberturas(doc: Document, datos: dict, assets: str, idx: int) -> None:
    """Render Section IV: Coverages matrix."""
    _titulo(doc, idx, "COBERTURAS")
    _tabla_matriz(doc, COBERTURAS, datos["aseguradoras"], "coberturas", assets)


def _seccion_deducibles(doc: Document, datos: dict, assets: str, idx: int) -> None:
    """Render Section V: Deductibles matrix."""
    _titulo(doc, idx, "DEDUCIBLES")
    _tabla_matriz(
        doc, DEDUCIBLES, datos["aseguradoras"], "deducibles", assets, ancho_concepto=2.1
    )


TEXTO_PREVALENCIA = [
    (
        "La información sobre el listado de Coberturas, Deducibles, Valores Asegurados, Cláusulas "
        "Adicionales, Sublímites, Primas, Tasas, Conclusiones, etc., es meramente ilustrativa y fue "
        "tomada de la cotización original y en firme presentada por el mercado asegurador. Los ítems "
        "mencionados se toman en forma parcial y a manera de ejemplo para facilitar la comparación, "
        "análisis y conclusión de lo que podría constituirse en la cotización más favorable al riesgo "
        "que se pretende trasladar."
    ),
    (
        "Multiriesgos de Colombia Ltda. expresamente hace extensivo al cliente los términos y "
        "condiciones oficiales de cada una de las Compañías de Seguros que demostraron su interés en el "
        "riesgo planteado y que presentaron su propuesta. En consecuencia, se deja establecido y "
        "expresamente pactado que, para todos los efectos técnicos y legales, prevalecen las Condiciones "
        "Técnicas y Económicas presentadas oficialmente por cada una de las Compañías de Seguros "
        "colombianas."
    ),
]


def _seccion_prevalencia(doc: Document, idx: int) -> None:
    """Render Section VI: Terms Prevalence disclaimer."""
    _titulo(doc, idx, "PREVALENCIA DE LOS TÉRMINOS COTIZADOS")
    tabla = doc.add_table(rows=1, cols=1)
    _bordes_tabla(tabla)
    celda = tabla.rows[0].cells[0]
    celda.text = ""
    for i, bloque in enumerate(TEXTO_PREVALENCIA):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(bloque)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = FUENTE


TEXTO_LEG2 = (
    "El seguro cubre las pérdidas o daños materiales que sean consecuencia de un defecto, "
    "pero excluye el costo de corregir el defecto original que dio lugar al daño. Desde el "
    "enfoque de suscripción y ajuste de siniestros se separa causa y consecuencia: la causa "
    "(defecto) NO está asegurada; la consecuencia (daño resultante) SÍ está asegurada."
)
TEXTO_LEG3 = (
    "El seguro cubre tanto los daños consecuenciales como el daño en el propio elemento "
    "defectuoso, excluyendo únicamente el costo adicional necesario para mejorar, rediseñar "
    "o corregir el defecto más allá del estándar original (improvement cost)."
)


def _anexos(doc: Document, datos: dict, assets: str, indice_inicial: int) -> None:
    """Render all annexes at the end of the document."""

    def titulo_anexo(n: int, etiqueta: str) -> None:
        doc.add_page_break()
        p = doc.add_paragraph()
        r = p.add_run(f"ANEXO {n}")
        r.bold = True
        r.underline = True
        r.font.size = Pt(12)
        r.font.name = FUENTE
        r.font.color.rgb = AZUL
        p2 = doc.add_paragraph()
        r2 = p2.add_run(etiqueta)
        r2.bold = True
        r2.underline = True
        r2.font.size = Pt(11)
        r2.font.name = FUENTE

    n = 0
    cfg = datos.get("anexos", {})

    if cfg.get("leg", True):
        n += 1
        titulo_anexo(n, "ALCANCE LEG 2 Y LEG 3")
        _parrafo(doc, "LEG 2 establece que:", negrita=True)
        _parrafo(doc, TEXTO_LEG2)
        _parrafo(doc, "LEG 3 establece que:", negrita=True)
        _parrafo(doc, TEXTO_LEG3)

    if cfg.get("tasa_prorroga", True):
        n += 1
        titulo_anexo(n, "TASAS DE PRÓRROGA")
        tabla = doc.add_table(rows=1, cols=2)
        _bordes_tabla(tabla)
        _texto(
            tabla.rows[0].cells[0],
            "COMPAÑÍA DE SEGUROS",
            negrita=True,
            tam=9,
            color=RGBColor(0xFF, 0xFF, 0xFF),
        )
        _sombrear(tabla.rows[0].cells[0], AZUL_HEX)
        _texto(
            tabla.rows[0].cells[1],
            "TASA DE PRÓRROGA",
            negrita=True,
            tam=9,
            color=RGBColor(0xFF, 0xFF, 0xFF),
        )
        _sombrear(tabla.rows[0].cells[1], AZUL_HEX)
        for ase in datos["aseguradoras"]:
            fila = tabla.add_row()
            fila.cells[0].width = Inches(2.2)
            fila.cells[1].width = Inches(4.3)
            _logo_en_celda(
                fila.cells[0],
                os.path.join(assets, "logos", ase.get("logo", "")),
                1.6,
            )
            _texto(
                fila.cells[1],
                ase.get("tasa_prorroga", "NO ESPECIFICA"),
                tam=9,
                centrado=False,
            )

    subj = [a for a in datos["aseguradoras"] if a.get("subjetividades")]
    if subj:
        n += 1
        titulo_anexo(n, "SUBJETIVIDADES Y CONDICIONES PARTICULARES POR ASEGURADORA")
        for ase in subj:
            _parrafo(doc, ase["nombre"], negrita=True, tam=10)
            for s in ase["subjetividades"]:
                _parrafo(doc, s, tam=9, vineta=True)

    obs = cfg.get("observaciones_adicionales") or []
    preg = cfg.get("preguntas_cliente") or []
    if obs or preg:
        n += 1
        titulo_anexo(n, "OBSERVACIONES Y PREGUNTAS DEL CLIENTE")
        for o in obs:
            _parrafo(doc, o, tam=9, vineta=True)
        for item in preg:
            _parrafo(doc, item.get("pregunta", ""), negrita=True, tam=10)
            _parrafo(doc, item.get("respuesta", ""), tam=9)


def _asegura_membrete(doc: Document, assets: str) -> None:
    """Ensure the document has the corporate letterhead in the header."""
    seccion = doc.sections[0]
    encabezado = seccion.header
    tiene_imagen = bool(encabezado._element.findall(".//" + qn("a:blip")))
    if tiene_imagen:
        return
    logo = os.path.join(assets, "logo_mrc.png")
    p = encabezado.paragraphs[0] if encabezado.paragraphs else encabezado.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(logo):
        p.add_run().add_picture(logo, width=Inches(2.2))


# --------------------------------------------------------------------------- #
# Public API                                                                   #
# --------------------------------------------------------------------------- #

def render_comparative(
    data: dict,
    output_path: str | None = None,
    assets_dir: str | None = None,
) -> str:
    """Generate the Phase II comparative Word document.

    Args:
        data: Consolidated data dict (datos.json structure) with keys:
              meta, aseguradoras, recomendacion, marcadores_peor, anexos.
        output_path: Where to save the .docx. Auto-generated if None.
        assets_dir: Path to the assets directory with logos and templates.

    Returns:
        Absolute path to the generated .docx file.
    """
    assets = assets_dir or str(settings.ASSETS_DIR)
    plantilla = os.path.join(assets, "plantilla_membrete.docx")

    if output_path is None:
        output_path = os.path.join(
            tempfile.gettempdir(),
            f"COMPARATIVO_FASE_2_{data.get('meta', {}).get('fecha', 'output').replace('/', '-')}.docx",
        )

    logger.info("Generando comparativo Fase II: %s", output_path)

    doc = Document(plantilla) if os.path.exists(plantilla) else Document()
    if not os.path.exists(plantilla):
        for s in doc.sections:
            s.left_margin = s.right_margin = Inches(1)
    _asegura_membrete(doc, assets)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(10)

    i = 0
    _seccion_general(doc, data.get("meta", {}), i)
    i += 1
    _seccion_cotizaciones(doc, data, assets, i)
    i += 1
    _seccion_recomendacion(doc, data, assets, i)
    i += 1
    doc.add_page_break()
    _seccion_coberturas(doc, data, assets, i)
    i += 1
    doc.add_page_break()
    _seccion_deducibles(doc, data, assets, i)
    i += 1
    _seccion_prevalencia(doc, i)
    i += 1
    _anexos(doc, data, assets, i)

    doc.save(output_path)
    logger.info("Comparativo Fase II generado: %s", output_path)
    return output_path
