# -*- coding: utf-8 -*-
"""Concept Mapper service — Dictionary and normalization matrix based on concept_matrix.xlsx.

Maps insurer-specific terms, variants, and synonyms to the official SFC
(Superintendencia Financiera de Colombia) 'Concepto Técnico Principal (SFC)'
according to insurance category ('Categoria').
"""
from __future__ import annotations

import logging
import os
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Optional

import openpyxl

from app.core.config import settings

logger = logging.getLogger(__name__)


def normalize_str(text: Any) -> str:
    """Normalize string by removing accents, special characters, and excess whitespace."""
    if text is None:
        return ""
    s = str(text)
    # NFKD normalization to remove accents/diacritics
    s = unicodedata.normalize("NFKD", s).encode("ASCII", "ignore").decode("utf-8")
    # Replace non-alphanumeric with spaces
    s = re.sub(r"[^\w\s]", " ", s.lower())
    # Collapse multiple spaces
    return re.sub(r"\s+", " ", s).strip()


BUILTIN_SYNONYMS: dict[str, dict[str, list[str]]] = {
    "Autos": {
        "RCE (Daños a bienes de Terceros, Lesiones o muerte a una persona, Lesiones o muerte a dos o más personas)": [
            "rce", "rce_danos", "rce_basico", "danos a terceros", "lesiones", "muerte", "responsabilidad civil extracontractual", "responsabilidad civil", "ded_rce"
        ],
        "Perdida parcial y total Daños": [
            "perdida_parcial_total_danos", "danos materiales", "perdida parcial por danos", "perdida total danos", "danos al vehiculo", "ded_perdida_total_danos", "ded_perdida_parcial_danos", "perdida total danos", "perdida parcial danos"
        ],
        "Pérdida parcial y total por Hurto": [
            "perdida_parcial_total_hurto", "hurto", "perdida parcial por hurto", "perdida total por hurto", "hurto calificado", "ded_perdida_total_hurto", "ded_perdida_parcial_hurto", "perdida total hurto", "perdida parcial hurto"
        ],
        "Protección Patrimonial": [
            "proteccion_patrimonial", "amparo patrimonial", "cobertura patrimonial"
        ],
        "Terremoto, temblor": [
            "terremoto_temblor", "eventos de la naturaleza", "terremoto", "temblor", "erupcion", "inundacion", "ded_terremoto"
        ],
        "Asistencia Jurídica": [
            "asistencia_juridica", "defensa legal", "asistencia penal", "asistencia legal"
        ],
        "Accidentes Personales": [
            "accidentes_personales", "accidentes ocupantes", "muerte accidental"
        ],
        "Gastos de Transporte Por Pérdidas Totales": [
            "gastos_transporte", "transporte por perdida total", "movilizacion"
        ],
        "Conductor Elegido": [
            "conductor_elegido", "servicio de conductor", "conductor profesional"
        ],
        "Vehiculo de Reemplazo": [
            "vehiculo_reemplazo", "carro de reemplazo", "vehiculo sustituto"
        ],
        "Otros amparos": [
            "otros_amparos", "asistencias", "grua", "asistencia en viaje", "llaves"
        ],
    },
    "Copropiedades": {
        "Incendio y/o rayo, Explosión, Daños por agua, Anegación": [
            "incendio_rayo_explosion_agua", "danos_materiales", "incendio", "modulo basico de incendio", "explosion", "danos por agua", "ded_incendio", "ded_explosion_agua"
        ],
        "Terremoto, temblor y/o erupción volcánica, HMACC, AMIT": [
            "terremoto_hmacc_amit", "terremoto", "hmacc", "amit", "sabotaje", "terrorismo", "ded_terremoto"
        ],
        "Rotura de vidrios": ["rotura_vidrios", "vidrios", "rotura accidental de vidrios"],
        "Pago cuotas de Administración": ["cuotas_administracion", "perdida de cuotas", "pago cuotas"],
        "Remoción de escombros": ["remocion_escombros", "escombros"],
        "Rotura de maquinaria (daño interno)": ["rotura_maquinaria", "dano interno", "maquinaria", "rotura de maquinaria", "ded_rotura_maquinaria"],
        "Sustracción con violencia": ["sustraccion_violencia", "hurto", "hurto calificado"],
        "Básico": ["rce_basico", "rce", "responsabilidad civil extracontractual", "rce predios labores y operaciones", "ded_rce"],
        "RCE Básico": ["rce_basico", "rce", "responsabilidad civil", "ded_rce"],
        "D&O": ["dno_directores", "directores", "administradores", "d&o", "dno"],
        "Manejo": ["manejo_fraude", "manejo", "fraude", "infidelidad", "ded_manejo"],
    },
    "Hogar": {
        "Incendio y riesgos aliados": ["incendio_edificio", "incendio_equipos", "incendio_muebles", "incendio_obras_arte", "incendio", "danos_materiales", "ded_incendio"],
        "Actos Mal Intencionados de Terceros": ["amit_edificio", "amit_equipos", "amit_muebles", "amit_obras_arte", "amit", "actos mal intencionados", "ded_amit"],
        "Terremoto, Temblor o Erupción": ["terremoto_edificio", "terremoto_equipos", "terremoto_muebles", "terremoto_obras_arte", "terremoto", "ded_terremoto"],
        "Hurto Calificado / Hurto Simple": ["hurto_equipos", "hurto_muebles", "hurto_obras_arte", "hurto", "sustraccion", "hurto calificado", "hurto simple", "ded_hurto"],
        "Equipos Móviles, Portátiles y Joyas": ["equipos_moviles_joyas", "equipos moviles", "joyas", "portatiles"],
        "Asistencias Domiciliarias": ["asistencias_domiciliarias", "asistencias", "plomeria", "electricidad", "cerrajeria"],
        "RCE Familiar": ["rce_familiar", "rce", "responsabilidad civil familiar", "cabeza de familia"],
    },
}


class ConceptEntry:
    """Represents a row in concept_matrix.xlsx."""

    def __init__(
        self,
        key: str,
        category: str,
        sfc_concept: str,
        insurer_synonyms: dict[str, str],
        all_synonyms: list[str],
    ) -> None:
        self.key = key
        self.category = category
        self.category_normalized = normalize_str(category)
        self.sfc_concept = sfc_concept
        self.sfc_normalized = normalize_str(sfc_concept)
        self.insurer_synonyms = insurer_synonyms  # normalized_insurer -> term
        self.all_synonyms = all_synonyms
        self.synonyms_normalized = [normalize_str(s) for s in all_synonyms if s]


class ConceptMapper:
    """Manages loading and fuzzy lookup for insurance concepts across insurers and categories."""

    def __init__(self, matrix_path: Optional[str | Path] = None) -> None:
        self.matrix_path = (
            Path(matrix_path)
            if matrix_path
            else Path(settings.ASSETS_DIR) / "concept_matrix.xlsx"
        )
        self.entries: list[ConceptEntry] = []
        self.insurers: list[str] = []
        self._load_matrix()

    def _load_matrix(self) -> None:
        """Load concept_matrix.xlsx and index all concepts and synonyms."""
        if not self.matrix_path.exists():
            logger.warning(
                "Matriz de conceptos no encontrada en: %s. Operando sin mapeo de conceptos.",
                self.matrix_path,
            )
            return

        try:
            wb = openpyxl.load_workbook(self.matrix_path, data_only=True)
            ws = wb.active

            header_row = [cell.value for cell in ws[1]]
            # Col A: Key, Col B: Categoria, Col C: Concepto Técnico Principal (SFC)
            # Col D onwards: Insurer names (e.g. Axa Colpatria, Berkley, etc.)
            raw_insurers = [str(h) for h in header_row[3:] if h is not None]
            self.insurers = [normalize_str(h) for h in raw_insurers]

            for row in ws.iter_rows(min_row=2, values_only=True):
                key = str(row[0] or "").strip()
                cat = str(row[1] or "").strip()
                sfc = str(row[2] or "").strip()

                if not sfc or not cat:
                    continue

                insurer_synonyms: dict[str, str] = {}
                all_syns: list[str] = []

                for idx, val in enumerate(row[3:]):
                    if val is not None and str(val).strip():
                        term = str(val).strip()
                        all_syns.append(term)
                        if idx < len(self.insurers):
                            insurer_synonyms[self.insurers[idx]] = term

                entry = ConceptEntry(
                    key=key,
                    category=cat,
                    sfc_concept=sfc,
                    insurer_synonyms=insurer_synonyms,
                    all_synonyms=all_syns,
                )
                self.entries.append(entry)

            logger.info(
                "ConceptMapper cargó exitosamente %d conceptos desde %s",
                len(self.entries),
                self.matrix_path,
            )
        except Exception as e:
            logger.exception("Error al cargar concept_matrix.xlsx: %s", e)

    @staticmethod
    def normalize_category(category: str) -> str:
        """Normalize category input into standard layout/system names.

        Supported standards:
        - 'Autos'
        - 'Copropiedades'
        - 'Hogar'
        - 'Todo_Riesgo_Construccion'
        - 'Pyme'
        - 'Maquinaria_Equipo'
        """
        c = normalize_str(category)
        if "auto" in c or "vehiculo" in c or "carro" in c:
            return "Autos"
        if "copropiedad" in c or "horizontal" in c or "edificio" in c or "conjunto" in c:
            return "Copropiedades"
        if "hogar" in c or "habitacion" in c or "casa" in c or "residencial" in c:
            return "Hogar"
        if (
            "construc" in c
            or "trc" in c
            or "obra" in c
            or "ingenieria" in c
            or "montaje" in c
        ):
            return "Todo_Riesgo_Construccion"
        if "pyme" in c or "empresa" in c:
            return "Pyme"
        if "maquinaria" in c or "equipo" in c:
            return "Maquinaria_Equipo"
        return category.strip() or "Todo_Riesgo_Construccion"

    def get_category_alias_keys(self, category: str) -> list[str]:
        """Return acceptable category values in the excel matrix matching the category."""
        target = self.normalize_category(category)
        if target == "Autos":
            return ["autos"]
        if target == "Copropiedades":
            return ["copropiedades"]
        if target == "Hogar":
            return ["hogar"]
        if target == "Todo_Riesgo_Construccion":
            return ["todo riesgo construccion", "trc"]
        if target == "Pyme":
            return ["pyme"]
        if target == "Maquinaria_Equipo":
            return ["maquinaria"]
        return [normalize_str(target)]

    def map_concept(
        self,
        term: str,
        category: str,
        insurer: Optional[str] = None,
        threshold: float = 0.70,
    ) -> Optional[str]:
        """Find the canonical 'Concepto Técnico Principal (SFC)' for a given term.

        Args:
            term: The concept or phrase found in a quote or extracted text.
            category: The insurance category (e.g. 'Autos', 'Hogar', 'Copropiedades').
            insurer: Optional insurer identifier/name (e.g. 'sura', 'zurich', 'chubb').
            threshold: Minimum fuzzy similarity ratio (0.0 to 1.0).

        Returns:
            The official 'Concepto Técnico Principal (SFC)' string, or None if no match.
        """
        n_term = normalize_str(term)
        if not n_term:
            return None

        cat_aliases = self.get_category_alias_keys(category)
        # We always also check 'todas' as global shared concepts
        cat_aliases.extend(["todas", "linea de negocio"])

        # Filter candidate entries for this category
        candidates = [
            e for e in self.entries if any(alias in e.category_normalized for alias in cat_aliases)
        ]
        if not candidates:
            candidates = self.entries

        # 1. Exact match on insurer-specific term if insurer is specified
        if insurer:
            n_ins = normalize_str(insurer)
            for e in candidates:
                for ins_name, ins_term in e.insurer_synonyms.items():
                    if n_ins in ins_name or ins_name in n_ins:
                        if normalize_str(ins_term) == n_term:
                            return e.sfc_concept

        # 2. Exact match against SFC concept or any synonym
        for e in candidates:
            if e.sfc_normalized == n_term:
                return e.sfc_concept
            for syn_norm in e.synonyms_normalized:
                if syn_norm == n_term:
                    return e.sfc_concept

        # 3. Substring match (e.g., if one contains the other)
        best_sub: Optional[str] = None
        best_sub_len = 0
        for e in candidates:
            if n_term in e.sfc_normalized or e.sfc_normalized in n_term:
                if len(e.sfc_normalized) > best_sub_len:
                    best_sub = e.sfc_concept
                    best_sub_len = len(e.sfc_normalized)
            for syn_norm in e.synonyms_normalized:
                if n_term in syn_norm or syn_norm in n_term:
                    if len(syn_norm) > best_sub_len:
                        best_sub = e.sfc_concept
                        best_sub_len = len(syn_norm)
        if best_sub and best_sub_len >= 4:
            return best_sub

        # 4. Fuzzy similarity matching using SequenceMatcher
        best_match: Optional[str] = None
        best_score = 0.0

        for e in candidates:
            score = SequenceMatcher(None, n_term, e.sfc_normalized).ratio()
            if score > best_score:
                best_score = score
                best_match = e.sfc_concept

            for syn_norm in e.synonyms_normalized:
                score = SequenceMatcher(None, n_term, syn_norm).ratio()
                if score > best_score:
                    best_score = score
                    best_match = e.sfc_concept

        if best_score >= threshold and best_match:
            return best_match

        # 5. Fallback against built-in domain dictionary for active categories
        norm_cat = self.normalize_category(category)
        cat_builtins = BUILTIN_SYNONYMS.get(norm_cat, {})
        for sfc_canon, aliases in cat_builtins.items():
            if n_term == normalize_str(sfc_canon):
                return sfc_canon
            for al in aliases:
                norm_al = normalize_str(al)
                if n_term == norm_al or n_term in norm_al or norm_al in n_term:
                    return sfc_canon

        return None

    def get_canonical_concepts_for_category(self, category: str) -> list[str]:
        """Get the ordered list of SFC concepts defined for a category."""
        cat_aliases = self.get_category_alias_keys(category)
        cat_aliases.extend(["todas"])
        seen = set()
        concepts = []
        for e in self.entries:
            if any(alias in e.category_normalized for alias in cat_aliases):
                if e.sfc_concept not in seen:
                    seen.add(e.sfc_concept)
                    concepts.append(e.sfc_concept)
        return concepts

    def normalize_extracted_dict(
        self,
        raw_dict: dict[str, Any],
        category: str,
        insurer: Optional[str] = None,
    ) -> dict[str, Any]:
        """Normalize keys in an extracted dictionary to their canonical SFC concepts.

        If a key cannot be mapped, it is retained as-is.
        """
        normalized: dict[str, Any] = {}
        for k, v in raw_dict.items():
            mapped = self.map_concept(k, category=category, insurer=insurer)
            key_to_use = mapped if mapped else k
            normalized[key_to_use] = v
        return normalized


# ---------------------------------------------------------------------------
# Singleton instance
# ---------------------------------------------------------------------------
_mapper: Optional[ConceptMapper] = None


def get_concept_mapper() -> ConceptMapper:
    """Get or initialize the cached ConceptMapper singleton instance."""
    global _mapper
    if _mapper is None:
        _mapper = ConceptMapper()
    return _mapper


def map_concept_to_sfc(
    term: str,
    category: str,
    insurer: Optional[str] = None,
    threshold: float = 0.70,
) -> Optional[str]:
    """Convenience helper to map any term to SFC canonical concept."""
    return get_concept_mapper().map_concept(
        term=term, category=category, insurer=insurer, threshold=threshold
    )
