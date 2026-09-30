# -*- coding: utf-8 -*-
"""Unit tests for DOCX Renderer service across category layouts."""
from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.services.docx_renderer import render_comparative


class TestDocxRendererLayouts(unittest.TestCase):
    """Test suite for Word document rendering using assets/layouts templates."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.sample_data = {
            "meta": {
                "tomador": "Test Corp S.A.S.",
                "identificacion": "900.555.444-3",
                "asegurado": "Test Corp S.A.S.",
                "beneficiario": "Terceros Afectados",
                "marca": "CHEVROLET",
                "placa": "XYZ123",
                "linea": "TRACKER",
                "modelo": "2024",
                "fecha": "28/09/2026",
                "valor_asegurado": 85000000,
                "ubicacion": "Bogotá D.C.",
                "descripcion_proyecto": "Riesgo estándar",
            },
            "aseguradoras": [
                {
                    "id": "sura",
                    "nombre": "Seguros SURA",
                    "logo": "sura.png",
                    "opciones": [
                        {
                            "etiqueta": "Opción 1",
                            "tasa": "1,2 ‰",
                            "prima": 2500000.0,
                            "modalidad": "Todo Riesgo",
                        }
                    ],
                    "coberturas": {
                        "RCE Daños a bienes de Terceros": 1000000000,
                        "Perdida parcial y total Daños": "100%",
                        "Incendio y/o rayo": "100%",
                    },
                    "deducibles": {
                        "RCE": "10% min 1 SMMLV",
                        "Incendio y/o rayo": "10% min 2 SMMLV",
                    },
                    "subjetividades": ["Inspección técnica"],
                },
                {
                    "id": "zurich",
                    "nombre": "Zurich Colombia",
                    "logo": "zurich.png",
                    "opciones": [
                        {
                            "etiqueta": "Opción 1",
                            "tasa": "1,1 ‰",
                            "prima": 2350000.0,
                            "modalidad": "Todo Riesgo",
                        }
                    ],
                    "coberturas": {
                        "RCE Daños a bienes de Terceros": 1200000000,
                        "Perdida parcial y total Daños": "100%",
                        "Incendio y/o rayo": "100%",
                    },
                    "deducibles": {
                        "RCE": "10% min 1 SMMLV",
                        "Incendio y/o rayo": "10% min 2 SMMLV",
                    },
                    "subjetividades": ["SARLAFT"],
                },
            ],
            "recomendacion": {
                "aseguradora_id": "zurich",
                "opcion": "Opción 1",
                "vinetas": ["Mayor cobertura económica", "Tarifa más competitiva"],
            },
        }

    def test_render_autos_layout(self):
        """Test rendering Autos layout."""
        out = os.path.join(self.temp_dir, "comparativo_autos.docx")
        path = render_comparative(self.sample_data, output_path=out, categoria="Autos")
        self.assertTrue(os.path.exists(path))
        doc = Document(path)
        self.assertGreater(len(doc.tables), 5)
        # Verify tomador filled in General Info table
        t1_text = " ".join(c.text for row in doc.tables[1].rows for c in row.cells)
        self.assertIn("Test Corp S.A.S.", t1_text)

    def test_render_copropiedades_layout(self):
        """Test rendering Copropiedades layout."""
        out = os.path.join(self.temp_dir, "comparativo_copropiedades.docx")
        path = render_comparative(self.sample_data, output_path=out, categoria="Copropiedades")
        self.assertTrue(os.path.exists(path))
        doc = Document(path)
        self.assertGreater(len(doc.tables), 4)

    def test_render_hogar_layout(self):
        """Test rendering Hogar layout."""
        out = os.path.join(self.temp_dir, "comparativo_hogar.docx")
        path = render_comparative(self.sample_data, output_path=out, categoria="Hogar")
        self.assertTrue(os.path.exists(path))
        doc = Document(path)
        self.assertGreater(len(doc.tables), 4)

    def test_render_trc_layout(self):
        """Test rendering Todo Riesgo Construcción layout."""
        out = os.path.join(self.temp_dir, "comparativo_trc.docx")
        path = render_comparative(self.sample_data, output_path=out, categoria="Todo_Riesgo_Construccion")
        self.assertTrue(os.path.exists(path))
        doc = Document(path)
        self.assertGreater(len(doc.tables), 4)


if __name__ == "__main__":
    unittest.main()
