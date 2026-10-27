# -*- coding: utf-8 -*-
"""Tests for the Management Reports (Informes de Gestión) module."""
from __future__ import annotations

from pathlib import Path
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.informe_gestion_parser import (
    format_colombian_nit,
    inspect_flokzu_xlsx,
    process_flokzu_xlsx,
)
from app.services.informe_gestion_renderer import render_informe_gestion

EXAMPLE_DIR = Path(r"c:\jair\duban\API\ejemplo informe gestion")
XLSX_PATH = EXAMPLE_DIR / "multiriesgos-CUMP-2026-10-22.xlsx"
REF_DOCX_PATH = EXAMPLE_DIR / "Informe_Gestion_SONACOL_2026-09.docx"


def test_format_colombian_nit():
    assert format_colombian_nit("8301292898") == "830.129.289-8"
    assert format_colombian_nit("830129289") == "830.129.289-8"
    assert format_colombian_nit("830.129.289-8") == "830.129.289-8"


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="Reference Excel file not present")
def test_sonacol_sept_2026_exact_match(tmp_path: Path):
    file_bytes = XLSX_PATH.read_bytes()

    # 1. Inspect
    insp = inspect_flokzu_xlsx(file_bytes)
    assert insp.ramo == "Seguros de cumplimiento"
    assert insp.fecha_corte == "2026-10-22"
    assert insp.total_filas == 1409
    assert any(p.periodo == "2026-09" for p in insp.periodos)
    sonacol_sug = next(c for c in insp.clientes_sugeridos if c.id == "SONACOL")
    assert sonacol_sug.solicitudes_por_periodo["2026-09"] == 13

    # 2. Process
    data = process_flokzu_xlsx(
        file_bytes=file_bytes,
        periodo="2026-09",
        cliente_nombre_corto="SONACOL",
        cliente_nombre_completo="SONACOL – Solarte Nacional de Construcciones S.A.S.",
        cliente_nit="830.129.289-8",
    )
    kpis = data["kpis"]
    assert kpis["tramitadas"] == 13
    assert kpis["emitidas"] == 6
    assert kpis["novedades"] == 6
    assert kpis["declinadas"] == 0
    assert kpis["pendientes"] == 1
    assert kpis["efectividad_str"] == "92,3 %"
    assert kpis["gestion_str"] == "92,3 %"
    assert kpis["tiempo_promedio_str"] == "11,2 días"
    assert data["tiempos_medidas"]["promedio"] == "11,2 días (270 horas)"
    assert data["tiempos_medidas"]["mediana"] == "11,0 días"
    assert data["tiempos_medidas"]["extremos"] == "3,6 / 21,8 días"

    # 3. Render & compare against reference docx
    out_docx = tmp_path / "Informe_Gestion_SONACOL_2026-09_test.docx"
    assets_dir = Path(__file__).resolve().parent.parent / "assets"
    render_informe_gestion(data, out_docx, assets_dir=assets_dir)
    assert out_docx.exists()

    gen_doc = Document(str(out_docx))
    ref_doc = Document(str(REF_DOCX_PATH))

    assert len(gen_doc.tables) == len(ref_doc.tables) == 9
    for t_idx in range(9):
        gen_t = gen_doc.tables[t_idx]
        ref_t = ref_doc.tables[t_idx]
        assert len(gen_t.rows) == len(ref_t.rows), f"Table {t_idx} row count mismatch"
        for r_idx in range(len(ref_t.rows)):
            gen_cells = [c.text.strip() for c in gen_t.rows[r_idx].cells]
            ref_cells = [c.text.strip() for c in ref_t.rows[r_idx].cells]
            assert gen_cells == ref_cells, f"Table {t_idx} Row {r_idx} mismatch: {gen_cells} != {ref_cells}"


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="Reference Excel file not present")
def test_informes_api_endpoints():
    client = TestClient(app)
    file_bytes = XLSX_PATH.read_bytes()

    # 1. Inspect endpoint
    resp_insp = client.post(
        "/api/v1/informes/inspect",
        files={"file": ("multiriesgos-CUMP-2026-10-22.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp_insp.status_code == 200
    insp_json = resp_insp.json()
    assert insp_json["fecha_corte"] == "2026-10-22"

    # 2. Create job endpoint
    resp_job = client.post(
        "/api/v1/informes/jobs",
        data={
            "periodo": "2026-09",
            "cliente_nombre_corto": "SONACOL",
            "cliente_nombre_completo": "SONACOL – Solarte Nacional de Construcciones S.A.S.",
            "cliente_nit": "830.129.289-8",
        },
        files={"file": ("multiriesgos-CUMP-2026-10-22.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp_job.status_code == 202
    job_id = resp_job.json()["job_id"]

    # 3. Check status
    resp_status = client.get(f"/api/v1/informes/jobs/{job_id}/status")
    assert resp_status.status_code == 200
    st_json = resp_status.json()
    assert st_json["status"] == "completed"
    assert st_json["resumen"]["kpis"]["tramitadas"] == 13

    # 4. Direct download
    resp_dl = client.get(f"/api/v1/informes/jobs/{job_id}/download")
    assert resp_dl.status_code == 200
    assert len(resp_dl.content) > 10000
